import csv
import gzip
import io
import json
import secrets
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import HttpResponse, HttpResponseForbidden, HttpResponseBadRequest, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook, load_workbook
from .forms import ContactForm
from .models import Contact, ContactEmail, ContactPhone, Industry, Source
from .services import process_webhook, push_email_to_unisender

HEADERS = ["Фамилия", "Имя", "Отчество", "Компания", "Отрасль", "Источник", "Должность", "E-mail", "Статус e-mail", "Телефоны", "Топ-менеджмент", "GR", "Комментарий"]

@login_required
def contact_list(request):
    qs = Contact.objects.select_related("industry").prefetch_related("emails", "phones")
    q = request.GET.get("q", "").strip()
    category = request.GET.get("category", "")
    industry_id = request.GET.get("industry", "")
    if q:
        qs = qs.filter(Q(surname__icontains=q) | Q(first_name__icontains=q) | Q(patronymic__icontains=q) | Q(company__icontains=q) | Q(position__icontains=q) | Q(emails__address__icontains=q) | Q(phones__number__icontains=q)).distinct()
    if category == "top":
        qs = qs.filter(top_management=True)
    if category == "gr":
        qs = qs.filter(gr=True)
    if industry_id:
        qs = qs.filter(industry_id=industry_id)
    return render(request, "contacts/list.html", {"contacts": qs, "q": q, "category": category, "industry_id": industry_id, "industries": Industry.objects.all(), "total": qs.count()})

@login_required
def contact_edit(request, pk=None):
    contact = get_object_or_404(Contact, pk=pk) if pk else Contact()
    is_new = contact.pk is None
    if request.method == "POST":
        form = ContactForm(request.POST, instance=contact)
        if form.is_valid():
            try:
                with transaction.atomic():
                    saved = form.save()
                    new_emails = set(form.cleaned_data["emails_text"])
                    existing = {e.address: e for e in saved.emails.all()}
                    ContactEmail.objects.filter(contact=saved, address__in=(set(existing) - new_emails)).delete()
                    for address in new_emails - set(existing):
                        email_obj = ContactEmail.objects.create(contact=saved, address=address)
                        transaction.on_commit(lambda e=email_obj: push_email_to_unisender(e))
                    phone_values = form.cleaned_data["phones_text"]
                    saved.phones.all().delete()
                    ContactPhone.objects.bulk_create([ContactPhone(contact=saved, number=n) for n in phone_values])
                messages.success(request, "Контакт сохранён.")
                return redirect("contact_detail", pk=saved.pk)
            except IntegrityError:
                form.add_error("emails_text", "Один из адресов уже добавлен в базу. Обновите страницу и проверьте e-mail.")
    else:
        form = ContactForm(instance=contact)
    return render(request, "contacts/form.html", {"form": form, "contact": contact, "is_new": is_new})

@login_required
def contact_detail(request, pk):
    contact = get_object_or_404(Contact.objects.select_related("industry").prefetch_related("emails", "phones", "history__user"), pk=pk)
    return render(request, "contacts/detail.html", {"contact": contact})

@login_required
def contact_delete(request, pk):
    contact = get_object_or_404(Contact, pk=pk)
    if request.method == "POST":
        contact.delete()
        messages.success(request, "Контакт удалён.")
        return redirect("contact_list")
    return render(request, "contacts/confirm_delete.html", {"contact": contact})

def export_rows(queryset, include_empty=False):
    rows = []
    for c in queryset.prefetch_related("emails", "phones").select_related("industry", "source"):
        emails = list(c.emails.all())
        if not emails and include_empty:
            emails = [None]
        for email in emails:
            rows.append([c.surname, c.first_name, c.patronymic, c.company, c.industry.name if c.industry else "", c.source.name if c.source else "", c.position,
                         email.address if email else "", email.mailing_status if email else "",
                         "; ".join(c.phones.values_list("number", flat=True)),
                         "Да" if c.top_management else "Нет", "Да" if c.gr else "Нет", c.comments])
    return rows

@login_required
def export_contacts(request):
    qs = Contact.objects.all()
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(Q(surname__icontains=q) | Q(first_name__icontains=q) | Q(patronymic__icontains=q) | Q(company__icontains=q) | Q(position__icontains=q) | Q(emails__address__icontains=q) | Q(phones__number__icontains=q)).distinct()
    if request.GET.get("category") == "top":
        qs = qs.filter(top_management=True)
    if request.GET.get("category") == "gr":
        qs = qs.filter(gr=True)
    if request.GET.get("industry"):
        qs = qs.filter(industry_id=request.GET["industry"])
    rows = export_rows(qs, request.GET.get("include_no_email") == "1")
    if request.GET.get("format") == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Контакты"
        ws.append(HEADERS)
        for row in rows:
            ws.append(row)
        output = io.BytesIO()
        wb.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="contacts.xlsx"'
        return response
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(HEADERS)
    writer.writerows(rows)
    response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="contacts.csv"'
    return response

def parse_upload(upload):
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(upload, read_only=True, data_only=True)
        values = list(workbook.active.values)
    else:
        raw = upload.read()
        text = None
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise ValueError("Не удалось определить кодировку CSV (поддерживаются UTF-8 и Windows-1251).")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t,")
        except csv.Error:
            dialect = csv.excel
        values = list(csv.reader(io.StringIO(text), dialect))
    if not values:
        return []
    headers = [str(x or "").strip().lower() for x in values[0]]
    aliases = {"фамилия":"surname", "surname":"surname", "имя":"first_name", "first name":"first_name", "first_name":"first_name", "отчество":"patronymic", "patronymic":"patronymic", "компания":"company", "company":"company", "отрасль":"industry", "industry":"industry", "источник":"source", "source":"source", "должность":"position", "position":"position", "e-mail":"emails", "email":"emails", "e-mail адрес":"emails", "emails":"emails", "телефон":"phones", "телефоны":"phones", "phone":"phones", "phones":"phones", "топ-менеджмент":"top", "top management":"top", "top_management":"top", "gr":"gr", "направление gr":"gr", "комментарий":"comments", "comments":"comments"}
    mapping = [aliases.get(h, "") for h in headers]
    rows = []
    for number, raw_row in enumerate(values[1:], 2):
        if not any(raw_row):
            continue
        item = {"row_num": number}
        for index, field in enumerate(mapping):
            if field and index < len(raw_row):
                item[field] = str(raw_row[index] or "").strip()
        item["emails"] = [x.strip().lower() for x in item.get("emails", "").replace(",", ";").replace("\n", ";").split(";") if x.strip()]
        item["phones"] = [x.strip() for x in item.get("phones", "").replace("\n", ";").split(";") if x.strip()]
        item["top"] = item.get("top", "").lower() in ("да", "yes", "1", "true", "x")
        item["gr"] = item.get("gr", "").lower() in ("да", "yes", "1", "true", "x")
        rows.append(item)
    return rows

@login_required
def import_contacts(request):
    if request.method == "POST" and request.POST.get("action") == "commit":
        rows = request.session.pop("import_rows", [])
        mode = request.POST.get("mode", "skip")
        created = updated = skipped = 0
        for item in rows:
            emails = item.get("emails", [])
            matches = list(Contact.objects.filter(emails__address__in=emails).distinct()) if emails else []
            if len(matches) > 1 or (item.get("conflict") and len(matches) != 1):
                skipped += 1
                continue
            existing = matches[0] if matches else None
            if existing and mode == "skip":
                skipped += 1
                continue
            contact = existing or Contact()
            for key in ("surname", "first_name", "patronymic", "company", "position", "comments"):
                if key in item:
                    setattr(contact, key, item[key])
            contact.top_management = item.get("top", False)
            contact.gr = item.get("gr", False)
            if item.get("industry"):
                industry, _ = Industry.objects.get_or_create(name=item["industry"])
                contact.industry = industry
            if item.get("source"):
                source, _ = Source.objects.get_or_create(name=item["source"])
                contact.source = source
            contact.save()
            if existing:
                updated += 1
            else:
                created += 1
            for address in emails:
                owner = ContactEmail.objects.filter(address__iexact=address).first()
                if owner and owner.contact_id != contact.pk:
                    continue
                email_obj, created_email = ContactEmail.objects.get_or_create(address=address, defaults={"contact": contact})
                if created_email:
                    transaction.on_commit(lambda e=email_obj: push_email_to_unisender(e))
            for number in item.get("phones", []):
                ContactPhone.objects.get_or_create(contact=contact, number=number)
        messages.success(request, f"Импорт завершён: создано {created}, обновлено {updated}, пропущено {skipped}.")
        return redirect("contact_list")
    if request.method == "POST" and request.FILES.get("file"):
        try:
            rows = parse_upload(request.FILES["file"])
            for item in rows:
                existing = list(Contact.objects.filter(emails__address__in=item.get("emails", [])).distinct())
                item["conflict"] = len(existing) > 1
                item["matched"] = str(existing[0]) if len(existing) == 1 else ""
            request.session["import_rows"] = rows
            return render(request, "contacts/import_preview.html", {"rows": rows, "count": len(rows)})
        except Exception as exc:
            messages.error(request, f"Не удалось прочитать файл: {exc}")
    return render(request, "contacts/import.html")

@login_required
def unisender_sync(request):
    if request.method != "POST":
        return HttpResponseBadRequest("POST required")
    if not settings.UNISENDER_API_KEY or not settings.UNISENDER_LIST_ID:
        messages.error(request, "UniSender не настроен: задайте UNISENDER_API_KEY и UNISENDER_LIST_ID в .env.")
        return redirect("contact_list")
    pending = ContactEmail.objects.filter(synced_at__isnull=True)
    ok = errors = 0
    for email_obj in pending.iterator():
        if push_email_to_unisender(email_obj):
            ok += 1
        else:
            errors += 1
    messages.success(request, f"UniSender: отправлено {ok}, ошибок {errors}.")
    return redirect("contact_list")

@csrf_exempt
def unisender_webhook(request, token):
    if not settings.UNISENDER_WEBHOOK_TOKEN or not secrets.compare_digest(token, settings.UNISENDER_WEBHOOK_TOKEN):
        return HttpResponseForbidden("Forbidden")
    if request.method != "POST":
        return HttpResponseBadRequest("POST required")
    try:
        body = request.body or b"{}"
        if request.headers.get("Content-Encoding", "").lower() == "gzip":
            body = gzip.decompress(body)
        payload = json.loads(body)
    except json.JSONDecodeError:
        return HttpResponseBadRequest("Invalid JSON")
    return JsonResponse({"ok": True, "updated": process_webhook(payload)})
